# -*- coding: utf-8 -*-
"""python-docx 适配器：生成 Word 招标/采购文档（奇美拉 V2 直装后由沙箱执行器调用）

被 sandbox_runner 以 module="_luxy_wrappers.docx_wrap", function="create" 调用。
"""
import os


def create(out_path, title="招标文件", subtitle="", sections=None, tables=None):
    """生成 docx：
    sections: [{"heading": "...", "lines": ["..."]}] 或 ["段落文本"]
    tables:   [{"headers": [...], "rows": [[...]], "caption": "..."}]
    """
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.add_heading(str(title), level=0)
    if subtitle:
        p = doc.add_paragraph(str(subtitle))
        p.runs[0].font.size = Pt(11)
    for sec in (sections or []):
        if isinstance(sec, dict):
            if sec.get("heading"):
                doc.add_heading(str(sec["heading"]), level=1)
            for line in (sec.get("lines") or []):
                doc.add_paragraph(str(line))
        else:
            doc.add_paragraph(str(sec))
    for tb in (tables or []):
        if tb.get("caption"):
            doc.add_paragraph(str(tb["caption"]))
        headers = [str(x) for x in (tb.get("headers") or [])]
        rows = tb.get("rows") or []
        if headers:
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = "Table Grid"
            for i, h in enumerate(headers):
                t.rows[0].cells[i].text = h
            for row in rows:
                cells = t.add_row().cells
                for i, v in enumerate(row[:len(headers)]):
                    cells[i].text = "" if v is None else str(v)
    out_path = str(out_path)
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    doc.save(out_path)
    return {"path": out_path, "bytes": os.path.getsize(out_path),
            "sections": len(sections or []), "tables": len(tables or [])}


def __self_test__():
    import docx
    return {"ok": True, "module": docx.__name__,
            "version": getattr(docx, "__version__", "")}
