# -*- coding: utf-8 -*-
"""
Word(.docx) 只读支持的测试。

要点：
  1. classify_upload —— 把 .docx / .doc / .pdf / .txt 分清楚（.doc 要给出"另存为 docx"的路径）
  2. extract_text_from_docx —— 段落与**表格**都要出来（论文/报告里表格常带关键数据）
  3. 落点约定：与 PDF 一样转成 .txt → 下游工作流/数据集/知识库无需改动
"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
try:
    import _guards  # noqa: F401
except Exception:
    pass

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import lora_manager as lm  # noqa: E402


class TestClassifyUpload(unittest.TestCase):

    def test_known_kinds(self):
        cases = {
            "a.pdf": "pdf", "a.PDF": "pdf",
            "a.txt": "text", "a.md": "text", "a.text": "text",
            "a.docx": "docx", "A.DOCX": "docx",
            "a.doc": "legacy_doc", "a.DOC": "legacy_doc",
        }
        for name, want in cases.items():
            self.assertEqual(lm.classify_upload(name), want, name)

    def test_rejects_others(self):
        for name in ("a.xlsx", "a.pptx", "a.zip", "a", ""):
            self.assertIsNone(lm.classify_upload(name), name)

    def test_doc_is_distinguished_from_docx(self):
        """必须区分 .doc 与 .docx —— 否则用户传 .doc 只会看到一句含糊的"不支持"。"""
        self.assertNotEqual(lm.classify_upload("x.doc"), lm.classify_upload("x.docx"))

    def test_docx_in_upload_allowlist(self):
        self.assertIn(".docx", lm.UPLOAD_LIST_EXTS)


class TestDocxExtraction(unittest.TestCase):

    def test_paragraphs_and_tables(self):
        try:
            import docx
        except Exception as e:
            self.skipTest(f"python-docx 未安装: {e}")
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "t.docx")
            doc = docx.Document()
            doc.add_paragraph("本文提出了一种新的方法，用于验证 Word 解析。")
            doc.add_paragraph("第二段：数据集规模为 450 万句对。")
            t = doc.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "模型"
            t.cell(0, 1).text = "BLEU"
            t.cell(1, 0).text = "Transformer"
            t.cell(1, 1).text = "28.4"
            doc.save(path)
            text = lm.extract_text_from_docx(path)
        self.assertIn("新的方法", text)
        self.assertIn("450 万句对", text)
        # 表格必须出来，且保留行列结构
        self.assertIn("[表 1]", text)
        self.assertIn("Transformer", text)
        self.assertIn("28.4", text)
        self.assertIn("|", text, "表格应按行用 | 连接，便于下游理解结构")

    def test_empty_document_returns_empty(self):
        try:
            import docx
        except Exception as e:
            self.skipTest(f"python-docx 未安装: {e}")
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "empty.docx")
            docx.Document().save(path)
            self.assertEqual(lm.extract_text_from_docx(path).strip(), "")

    def test_bad_file_raises(self):
        """传个假 docx 应报错，而不是静默返回空串（否则会出现"上传成功但没内容"）。"""
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "fake.docx")
            with open(path, "wb") as f:
                f.write(b"not a real docx")
            with self.assertRaises(Exception):
                lm.extract_text_from_docx(path)


if __name__ == "__main__":
    unittest.main()
